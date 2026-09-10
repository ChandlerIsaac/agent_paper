"""Document loading utilities."""

import hashlib
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from langchain_core.documents import Document


HASH_BLOCK_SIZE = 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


def calculate_file_hash(path: Path) -> str:
    """Calculate a stable SHA-256 identifier without loading the whole file."""

    digest = hashlib.sha256()
    with path.open("rb") as file:
        while block := file.read(HASH_BLOCK_SIZE):
            digest.update(block)
    return digest.hexdigest()


def validate_document_path(file_path: str | Path) -> Path:
    """Validate and normalize a supported document path."""

    path = Path(file_path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"文档不存在：{path}")
    if not path.is_file():
        raise ValueError(f"输入路径不是文件：{path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        extensions = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"不支持的文件类型：{path.suffix}；支持：{extensions}")
    return path


def _metadata(path: Path, document_id: str, page: int, total_pages: int) -> dict:
    return {
        "document_id": document_id,
        "source": path.name,
        "page": page,
        "page_index": page - 1,
        "total_pages": total_pages,
        "file_type": path.suffix.lower().lstrip("."),
    }


def load_pdf_pages(file_path: str | Path) -> list[Document]:
    """Load a PDF as one LangChain document per page."""

    path = validate_document_path(file_path)
    if path.suffix.lower() != ".pdf":
        raise ValueError(f"文件不是 PDF：{path.name}")
    document_id = calculate_file_hash(path)

    try:
        with pymupdf.open(path) as pdf:
            if pdf.needs_pass:
                raise PermissionError(f"PDF 已加密，需要密码：{path.name}")
            if pdf.page_count == 0:
                raise ValueError(f"PDF 没有任何页面：{path.name}")

            documents = [
                Document(
                    page_content=page.get_text("text", sort=True).replace("\x00", "").strip(),
                    metadata=_metadata(path, document_id, index + 1, pdf.page_count),
                )
                for index, page in enumerate(pdf)
            ]
    except pymupdf.FileDataError as error:
        raise ValueError(f"PDF 文件损坏或格式无效：{path.name}") from error

    if not any(document.page_content for document in documents):
        raise ValueError(f"PDF 无可提取文字，可能是扫描件：{path.name}")
    return documents


def load_docx(file_path: str | Path) -> list[Document]:
    """Load a DOCX file as a single document."""

    path = validate_document_path(file_path)
    document_id = calculate_file_hash(path)
    docx = DocxDocument(path)
    text = "\n".join(p.text.strip() for p in docx.paragraphs if p.text.strip())
    if not text:
        raise ValueError(f"DOCX 没有可提取文字：{path.name}")
    return [Document(page_content=text, metadata=_metadata(path, document_id, 1, 1))]


def load_text(file_path: str | Path) -> list[Document]:
    """Load a UTF-8 Markdown or text file."""

    path = validate_document_path(file_path)
    text = path.read_text(encoding="utf-8").replace("\x00", "").strip()
    if not text:
        raise ValueError(f"文档内容为空：{path.name}")
    document_id = calculate_file_hash(path)
    return [Document(page_content=text, metadata=_metadata(path, document_id, 1, 1))]


def load_document(file_path: str | Path) -> list[Document]:
    """Dispatch a supported file to its format-specific loader."""

    path = validate_document_path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf_pages(path)
    if suffix == ".docx":
        return load_docx(path)
    return load_text(path)
